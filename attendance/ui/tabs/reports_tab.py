from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
    QTableWidget, QTableWidgetItem, QHeaderView,
    QLabel, QComboBox, QSpinBox, QFileDialog,
    QMessageBox, QStyle
)
from PySide6.QtCore import Qt
from sqlalchemy.orm import Session
from datetime import date
import calendar


_MONTHS = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
]

_COLS = [
    "Empleado", "Días con registro", "Días OK",
    "Días incompleto", "Días justificado",
    "Min. tardanza", "Min. horas extra", "Horas trabajadas"
]


class ReportsTab(QWidget):
    def __init__(self, session: Session):
        super().__init__()
        self.session = session
        self._records = []
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self)

        # ── Filter bar ────────────────────────────────────────
        filter_bar = QHBoxLayout()

        filter_bar.addWidget(QLabel("Mes:"))
        self.cb_month = QComboBox()
        for m in _MONTHS:
            self.cb_month.addItem(m)
        today = date.today()
        self.cb_month.setCurrentIndex(today.month - 1)
        filter_bar.addWidget(self.cb_month)

        filter_bar.addWidget(QLabel("Año:"))
        self.sp_year = QSpinBox()
        self.sp_year.setRange(2000, 2100)
        self.sp_year.setValue(today.year)
        filter_bar.addWidget(self.sp_year)

        filter_bar.addWidget(QLabel("Empleado:"))
        self.cb_employee = QComboBox()
        self.cb_employee.setEditable(True)
        self.cb_employee.addItem("Todos")
        filter_bar.addWidget(self.cb_employee)

        self.btn_generate = QPushButton("Generar Reporte")
        self.btn_generate.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_FileDialogContentsView))
        self.btn_generate.clicked.connect(self._generate)
        filter_bar.addWidget(self.btn_generate)

        self.btn_export = QPushButton("Exportar Excel")
        self.btn_export.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_DriveHDIcon))
        self.btn_export.setEnabled(False)
        self.btn_export.clicked.connect(self._export_excel)
        filter_bar.addWidget(self.btn_export)

        self.btn_export_pdf = QPushButton("Exportar PDF")
        self.btn_export_pdf.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_FileIcon))
        self.btn_export_pdf.setEnabled(False)
        self.btn_export_pdf.clicked.connect(self._export_pdf)
        filter_bar.addWidget(self.btn_export_pdf)

        filter_bar.addStretch()
        layout.addLayout(filter_bar)

        # ── Results table ─────────────────────────────────────
        self.table = QTableWidget()
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setColumnCount(len(_COLS))
        self.table.setHorizontalHeaderLabels(_COLS)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        layout.addWidget(self.table)

    def load_data(self):
        from models.employee import User
        self.cb_employee.clear()
        self.cb_employee.addItem("Todos")
        employees = self.session.query(User).filter(User.is_active == True).order_by(User.last_name).all()
        for emp in employees:
            name = f"{emp.last_name or ''}, {emp.first_name or ''}".strip(", ")
            self.cb_employee.addItem(f"{emp.identifier} — {name}", emp.id)

    def _generate(self):
        from models.employee import User
        from models.processed_attendance import ProcessedAttendance

        month = self.cb_month.currentIndex() + 1
        year = self.sp_year.value()
        _, last_day = calendar.monthrange(year, month)
        d_start = date(year, month, 1)
        d_end = date(year, month, last_day)

        emp_filter_id = self.cb_employee.currentData()

        query = self.session.query(ProcessedAttendance, User).join(
            User, ProcessedAttendance.employee_id == User.id
        ).filter(
            ProcessedAttendance.date >= d_start,
            ProcessedAttendance.date <= d_end
        )
        if emp_filter_id:
            query = query.filter(User.id == emp_filter_id)

        rows = query.order_by(User.last_name, User.first_name, ProcessedAttendance.date).all()

        # Aggregate by employee
        from collections import defaultdict
        agg = defaultdict(lambda: {
            "name": "", "total": 0, "ok": 0, "inc": 0, "jus": 0,
            "tard": 0, "ot": 0, "hours": 0.0
        })
        for pa, emp in rows:
            eid = emp.id
            agg[eid]["name"] = f"{emp.last_name or ''}, {emp.first_name or ''}".strip(", ") or emp.identifier
            agg[eid]["total"] += 1
            agg[eid]["ok"] += pa.status == "OK"
            agg[eid]["inc"] += pa.status == "INCOMPLETO"
            agg[eid]["jus"] += pa.status == "JUSTIFICADO"
            agg[eid]["tard"] += pa.tardiness_minutes or 0
            agg[eid]["ot"] += pa.overtime_minutes or 0
            agg[eid]["hours"] += pa.total_hours or 0.0

        self._records = list(agg.values())

        self.table.clearContents()
        self.table.setRowCount(len(self._records))
        for row, rec in enumerate(self._records):
            self.table.setItem(row, 0, QTableWidgetItem(rec["name"]))
            self.table.setItem(row, 1, QTableWidgetItem(str(rec["total"])))
            self.table.setItem(row, 2, QTableWidgetItem(str(rec["ok"])))
            self.table.setItem(row, 3, QTableWidgetItem(str(rec["inc"])))
            self.table.setItem(row, 4, QTableWidgetItem(str(rec["jus"])))
            self.table.setItem(row, 5, QTableWidgetItem(str(rec["tard"])))
            self.table.setItem(row, 6, QTableWidgetItem(str(rec["ot"])))
            self.table.setItem(row, 7, QTableWidgetItem(f"{rec['hours']:.2f}"))

        self.btn_export.setEnabled(bool(self._records))
        self.btn_export_pdf.setEnabled(bool(self._records))

        if not self._records:
            QMessageBox.information(self, "Sin datos", "No hay registros para el período seleccionado.")

    def _export_excel(self):
        if not self._records:
            return

        path, _ = QFileDialog.getSaveFileName(self, "Guardar Excel", "", "Excel Files (*.xlsx)")
        if not path:
            return

        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter

        month = self.cb_month.currentText()
        year = self.sp_year.value()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Reporte {month} {year}"

        header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        bold = Font(bold=True)
        center = Alignment(horizontal="center", vertical="center")

        # Title
        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = f"Reporte de Asistencia — {month} {year}"
        title_cell.font = Font(bold=True, size=13)
        title_cell.alignment = center
        ws.row_dimensions[1].height = 22

        # Headers
        for col, header in enumerate(_COLS, 1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = bold
            cell.fill = header_fill
            cell.alignment = center

        # Data
        for row_idx, rec in enumerate(self._records, 3):
            values = [
                rec["name"], rec["total"], rec["ok"], rec["inc"],
                rec["jus"], rec["tard"], rec["ot"], f"{rec['hours']:.2f}"
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = Alignment(vertical="top")

        # Column widths
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(14, max_len + 2)

        wb.save(path)
        QMessageBox.information(self, "Exportado", f"Reporte guardado en:\n{path}")

    def _export_pdf(self):
        if not self._records:
            return

        path, _ = QFileDialog.getSaveFileName(self, "Guardar PDF", "", "PDF Files (*.pdf)")
        if not path:
            return

        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        month = self.cb_month.currentText()
        year = self.sp_year.value()

        doc = SimpleDocTemplate(path, pagesize=landscape(A4),
                                leftMargin=1.5*cm, rightMargin=1.5*cm,
                                topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph(f"Reporte de Asistencia — {month} {year}", styles["Title"]))
        elements.append(Spacer(1, 0.4*cm))

        headers = [
            "Empleado", "Días\nregistro", "Días\nOK",
            "Días\nincompleto", "Días\njustificado",
            "Min.\ntardanza", "Min.\nhoras extra", "Horas\ntrabajadas"
        ]
        table_data = [headers]
        for rec in self._records:
            table_data.append([
                rec["name"],
                str(rec["total"]),
                str(rec["ok"]),
                str(rec["inc"]),
                str(rec["jus"]),
                str(rec["tard"]),
                str(rec["ot"]),
                f"{rec['hours']:.2f}",
            ])

        col_widths = [6*cm, 2*cm, 2*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm]
        tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(tbl)

        doc.build(elements)
        QMessageBox.information(self, "Exportado", f"Reporte PDF guardado en:\n{path}")
