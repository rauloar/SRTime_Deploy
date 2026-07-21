import io
from typing import List, Optional

from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def build_excel_workbook(
    sheet_name: str,
    headers: List[str],
    rows: List[list],
    col_widths: Optional[List[int]] = None,
    title: Optional[str] = None,
) -> Workbook:
    """Construye un Workbook de openpyxl con headers estilizados y datos."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    start_row = 1
    if title:
        ws.merge_cells(f"A1:{get_column_letter(len(headers))}1")
        title_cell = ws["A1"]
        title_cell.value = title
        title_cell.font = Font(bold=True, size=13)
        title_cell.alignment = CENTER
        ws.row_dimensions[1].height = 22
        start_row = 2

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = CENTER

    for row_idx, row_data in enumerate(rows, start_row + 1):
        for col_idx, val in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    if col_widths:
        for col_idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    else:
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(12, max_len + 2)

    return wb


def streaming_excel_response(
    wb: Workbook,
    filename: str,
) -> StreamingResponse:
    """Envuelve un Workbook en StreamingResponse para FastAPI."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
