from typing import List, Optional
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import date
import calendar
import io

from core.logging_config import get_logger
logger = get_logger(__name__)

from api.deps import get_db, get_current_user
from models.user import AuthUser

router = APIRouter()

_MONTH_NAMES = [
    "", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
]


@router.get("/monthly")
def monthly_summary(
    year: int = Query(...),
    month: int = Query(..., ge=1, le=12),
    employee_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    from models.employee import User
    from models.processed_attendance import ProcessedAttendance
    from collections import defaultdict

    _, last_day = calendar.monthrange(year, month)
    d_start = date(year, month, 1)
    d_end = date(year, month, last_day)

    query = db.query(ProcessedAttendance, User).join(
        User, ProcessedAttendance.employee_id == User.id
    ).filter(
        ProcessedAttendance.date >= d_start,
        ProcessedAttendance.date <= d_end
    )
    if employee_id:
        query = query.filter(User.id == employee_id)

    rows = query.all()

    agg = defaultdict(lambda: {
        "employee_id": None, "identifier": "", "name": "",
        "total": 0, "ok": 0, "incomplete": 0, "justified": 0,
        "tardiness_minutes": 0, "overtime_minutes": 0, "total_hours": 0.0
    })
    for pa, emp in rows:
        d = agg[emp.id]
        d["employee_id"] = emp.id
        d["identifier"] = emp.identifier
        d["name"] = f"{emp.last_name or ''}, {emp.first_name or ''}".strip(", ") or emp.identifier
        d["total"] += 1
        d["ok"] += pa.status == "OK"
        d["incomplete"] += pa.status == "INCOMPLETO"
        d["justified"] += pa.status == "JUSTIFICADO"
        d["tardiness_minutes"] += pa.tardiness_minutes or 0
        d["overtime_minutes"] += pa.overtime_minutes or 0
        d["total_hours"] += pa.total_hours or 0.0

    result = sorted(agg.values(), key=lambda x: x["name"])
    for r in result:
        r["total_hours"] = round(r["total_hours"], 2)

    return {
        "year": year,
        "month": month,
        "month_name": _MONTH_NAMES[month],
        "employees": result,
    }


@router.get("/monthly/export")
def export_monthly(
    year: int = Query(...),
    month: int = Query(..., ge=1, le=12),
    employee_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from models.employee import User
    from models.processed_attendance import ProcessedAttendance
    from collections import defaultdict

    _, last_day = calendar.monthrange(year, month)
    d_start = date(year, month, 1)
    d_end = date(year, month, last_day)

    query = db.query(ProcessedAttendance, User).join(
        User, ProcessedAttendance.employee_id == User.id
    ).filter(
        ProcessedAttendance.date >= d_start,
        ProcessedAttendance.date <= d_end
    )
    if employee_id:
        query = query.filter(User.id == employee_id)

    rows = query.all()
    agg = defaultdict(lambda: {
        "name": "", "total": 0, "ok": 0, "incomplete": 0, "justified": 0,
        "tard": 0, "ot": 0, "hours": 0.0
    })
    for pa, emp in rows:
        d = agg[emp.id]
        d["name"] = f"{emp.last_name or ''}, {emp.first_name or ''}".strip(", ") or emp.identifier
        d["total"] += 1
        d["ok"] += pa.status == "OK"
        d["incomplete"] += pa.status == "INCOMPLETO"
        d["justified"] += pa.status == "JUSTIFICADO"
        d["tard"] += pa.tardiness_minutes or 0
        d["ot"] += pa.overtime_minutes or 0
        d["hours"] += pa.total_hours or 0.0

    records = sorted(agg.values(), key=lambda x: x["name"])

    wb = openpyxl.Workbook()
    ws = wb.active
    month_label = _MONTH_NAMES[month]
    ws.title = f"Reporte {month_label} {year}"

    headers = [
        "Empleado", "Días con registro", "Días OK",
        "Días incompleto", "Días justificado",
        "Min. tardanza", "Min. horas extra", "Horas trabajadas"
    ]
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"A1:{get_column_letter(len(headers))}1")
    title_cell = ws["A1"]
    title_cell.value = f"Reporte de Asistencia — {month_label} {year}"
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = center
    ws.row_dimensions[1].height = 22

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = center

    for row_idx, rec in enumerate(records, 3):
        values = [
            rec["name"], rec["total"], rec["ok"], rec["incomplete"],
            rec["justified"], rec["tard"], rec["ot"], round(rec["hours"], 2)
        ]
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(14, max_len + 2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"reporte_{month_label.lower()}_{year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/monthly/export/pdf")
def export_monthly_pdf(
    year: int = Query(...),
    month: int = Query(..., ge=1, le=12),
    employee_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from models.employee import User
    from models.processed_attendance import ProcessedAttendance
    from collections import defaultdict

    _, last_day = calendar.monthrange(year, month)
    d_start = date(year, month, 1)
    d_end = date(year, month, last_day)

    query = db.query(ProcessedAttendance, User).join(
        User, ProcessedAttendance.employee_id == User.id
    ).filter(
        ProcessedAttendance.date >= d_start,
        ProcessedAttendance.date <= d_end
    )
    if employee_id:
        query = query.filter(User.id == employee_id)

    rows = query.all()
    agg = defaultdict(lambda: {
        "name": "", "total": 0, "ok": 0, "incomplete": 0, "justified": 0,
        "tard": 0, "ot": 0, "hours": 0.0
    })
    for pa, emp in rows:
        d = agg[emp.id]
        d["name"] = f"{emp.last_name or ''}, {emp.first_name or ''}".strip(", ") or emp.identifier
        d["total"] += 1
        d["ok"] += pa.status == "OK"
        d["incomplete"] += pa.status == "INCOMPLETO"
        d["justified"] += pa.status == "JUSTIFICADO"
        d["tard"] += pa.tardiness_minutes or 0
        d["ot"] += pa.overtime_minutes or 0
        d["hours"] += pa.total_hours or 0.0

    records = sorted(agg.values(), key=lambda x: x["name"])
    month_label = _MONTH_NAMES[month]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"Reporte de Asistencia — {month_label} {year}", styles["Title"]))
    elements.append(Spacer(1, 0.4*cm))

    headers = [
        "Empleado", "Días\nregistro", "Días\nOK",
        "Días\nincompleto", "Días\njustificado",
        "Min.\ntardanza", "Min.\nhoras extra", "Horas\ntrabajadas"
    ]
    table_data = [headers]
    for rec in records:
        table_data.append([
            rec["name"],
            str(rec["total"]),
            str(rec["ok"]),
            str(rec["incomplete"]),
            str(rec["justified"]),
            str(rec["tard"]),
            str(rec["ot"]),
            f"{rec['hours']:.2f}",
        ])

    col_widths = [6*cm, 2*cm, 2*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm]
    tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(tbl)

    doc.build(elements)
    buf.seek(0)

    filename = f"reporte_{month_label.lower()}_{year}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
