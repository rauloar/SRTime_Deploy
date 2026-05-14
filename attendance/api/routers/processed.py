from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
from datetime import date, datetime
import io

from core.logging_config import get_logger
logger = get_logger(__name__)

from api.deps import get_db, get_current_user, require_root, require_supervisor_or_root
from api.schemas import ProcessedAttendanceResponse
from models.processed_attendance import ProcessedAttendance
from models.employee import User as EmployeeModel
from models.shift import Shift
from models.user import AuthUser
from services.engine import AttendanceEngine

router = APIRouter()

class JustifyRequest(BaseModel):
    first_in: str
    last_out: str
    justification: str

class RequestJustifyRequest(BaseModel):
    first_in: str
    last_out: str
    justification: str


@router.post("/process-pending")
def process_pending(
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    engine = AttendanceEngine(db)
    count = engine.process_all(progress_callback=None, force_all=False)
    logger.info(f"Processed {count} pending days", extra={"user": current_user.username, "action": "process_pending"})
    return {"ok": True, "processed_days": count, "mode": "pending"}


@router.post("/reprocess-all")
def reprocess_all(
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    engine = AttendanceEngine(db)
    count = engine.process_all(progress_callback=None, force_all=True)
    logger.info(f"Reprocessed all: {count} days", extra={"user": current_user.username, "action": "reprocess_all"})
    return {"ok": True, "processed_days": count, "mode": "all"}

@router.get("/", response_model=List[ProcessedAttendanceResponse])
def read_processed(
    date_start: Optional[date] = Query(None),
    date_end: Optional[date] = Query(None),
    employee_id: Optional[int] = Query(None),
    skip: int = 0,
    limit: int = 500,
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    query = db.query(ProcessedAttendance)
    if date_start:
        query = query.filter(ProcessedAttendance.date >= date_start)
    if date_end:
        query = query.filter(ProcessedAttendance.date <= date_end)
    if employee_id:
        query = query.filter(ProcessedAttendance.employee_id == employee_id)
    return query.order_by(ProcessedAttendance.date.desc()).offset(skip).limit(limit).all()

@router.patch("/{record_id}/justify", response_model=ProcessedAttendanceResponse)
def justify_record(
    record_id: int,
    body: JustifyRequest,
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(require_root)
):
    rec = db.query(ProcessedAttendance).filter(ProcessedAttendance.id == record_id).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Record not found")
    if not body.justification.strip():
        raise HTTPException(status_code=400, detail="Justification text is required")

    from datetime import time as Time
    h_in, m_in = map(int, body.first_in.split(":"))
    h_out, m_out = map(int, body.last_out.split(":"))
    rec.first_in  = Time(h_in, m_in)
    rec.last_out  = Time(h_out, m_out)

    dt_in  = datetime.combine(rec.date, rec.first_in)
    dt_out = datetime.combine(rec.date, rec.last_out)
    rec.total_hours = round((dt_out - dt_in).total_seconds() / 3600.0, 2) if dt_out > dt_in else 0.0

    # Recalculate early departure and overtime
    employee = db.query(EmployeeModel).filter(EmployeeModel.id == rec.employee_id).first()
    if employee and employee.shift_id:
        shift = db.query(Shift).filter(Shift.id == employee.shift_id).first()
        if shift:
            exp_out = datetime.combine(rec.date, shift.expected_out)
            if dt_out < exp_out:
                rec.early_departure_minutes = int((exp_out - dt_out).total_seconds() / 60)
                rec.overtime_minutes = 0
            elif dt_out > exp_out:
                rec.overtime_minutes = int((dt_out - exp_out).total_seconds() / 60)
                rec.early_departure_minutes = 0
            else:
                rec.early_departure_minutes = 0
                rec.overtime_minutes = 0

    rec.status        = "JUSTIFICADO"
    rec.justification = body.justification
    rec.pending_first_in = None
    rec.pending_last_out = None
    rec.pending_justification = None
    rec.requested_by = None
    rec.previous_status = None
    db.commit()
    db.refresh(rec)
    logger.info(f"Record {record_id} justified", extra={"user": current_user.username, "action": "justify", "detail": str(record_id)})
    return rec


@router.patch("/{record_id}/request-justify", response_model=ProcessedAttendanceResponse)
def request_justify(
    record_id: int,
    body: RequestJustifyRequest,
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(require_supervisor_or_root),
):
    rec = db.query(ProcessedAttendance).filter(ProcessedAttendance.id == record_id).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Record not found")
    if rec.status == "PENDIENTE":
        raise HTTPException(status_code=400, detail="Already has a pending justification request")
    if not body.justification.strip():
        raise HTTPException(status_code=400, detail="Justification text is required")

    from datetime import time as Time
    h_in, m_in = map(int, body.first_in.split(":"))
    h_out, m_out = map(int, body.last_out.split(":"))

    rec.previous_status = rec.status
    rec.pending_first_in = Time(h_in, m_in)
    rec.pending_last_out = Time(h_out, m_out)
    rec.pending_justification = body.justification.strip()
    rec.requested_by = current_user.username
    rec.status = "PENDIENTE"

    db.commit()
    db.refresh(rec)
    logger.info(f"Record {record_id} justification requested", extra={"user": current_user.username, "action": "request_justify"})
    return rec


@router.patch("/{record_id}/approve-justify", response_model=ProcessedAttendanceResponse)
def approve_justify(
    record_id: int,
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(require_root),
):
    rec = db.query(ProcessedAttendance).filter(ProcessedAttendance.id == record_id).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Record not found")
    if rec.status != "PENDIENTE":
        raise HTTPException(status_code=400, detail="Record is not pending approval")

    rec.first_in = rec.pending_first_in
    rec.last_out = rec.pending_last_out

    dt_in = datetime.combine(rec.date, rec.first_in) if rec.first_in else None
    dt_out = datetime.combine(rec.date, rec.last_out) if rec.last_out else None
    rec.total_hours = round((dt_out - dt_in).total_seconds() / 3600.0, 2) if dt_in and dt_out and dt_out > dt_in else 0.0

    employee = db.query(EmployeeModel).filter(EmployeeModel.id == rec.employee_id).first()
    if employee and employee.shift_id and dt_out:
        shift = db.query(Shift).filter(Shift.id == employee.shift_id).first()
        if shift:
            exp_out = datetime.combine(rec.date, shift.expected_out)
            if dt_out < exp_out:
                rec.early_departure_minutes = int((exp_out - dt_out).total_seconds() / 60)
                rec.overtime_minutes = 0
            elif dt_out > exp_out:
                rec.overtime_minutes = int((dt_out - exp_out).total_seconds() / 60)
                rec.early_departure_minutes = 0
            else:
                rec.early_departure_minutes = 0
                rec.overtime_minutes = 0

    rec.status = "JUSTIFICADO"
    rec.justification = rec.pending_justification
    rec.pending_first_in = None
    rec.pending_last_out = None
    rec.pending_justification = None
    rec.requested_by = None
    rec.previous_status = None

    db.commit()
    db.refresh(rec)
    logger.info(f"Record {record_id} justify approved", extra={"user": current_user.username, "action": "approve_justify"})
    return rec


@router.patch("/{record_id}/reject-justify", response_model=ProcessedAttendanceResponse)
def reject_justify(
    record_id: int,
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(require_root),
):
    rec = db.query(ProcessedAttendance).filter(ProcessedAttendance.id == record_id).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Record not found")
    if rec.status != "PENDIENTE":
        raise HTTPException(status_code=400, detail="Record is not pending approval")

    rec.status = rec.previous_status or "INCOMPLETO"
    rec.pending_first_in = None
    rec.pending_last_out = None
    rec.pending_justification = None
    rec.requested_by = None
    rec.previous_status = None

    db.commit()
    db.refresh(rec)
    logger.info(f"Record {record_id} justify rejected", extra={"user": current_user.username, "action": "reject_justify"})
    return rec


@router.get("/summary")
def get_summary(
    date_start: Optional[date] = Query(None),
    date_end: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    query = db.query(ProcessedAttendance)
    if date_start:
        query = query.filter(ProcessedAttendance.date >= date_start)
    if date_end:
        query = query.filter(ProcessedAttendance.date <= date_end)
    records = query.all()

    by_status = {"OK": 0, "INCOMPLETO": 0, "JUSTIFICADO": 0}
    total_tard = total_ot = 0
    emp_with_tard = set()

    for r in records:
        by_status[r.status] = by_status.get(r.status, 0) + 1
        total_tard += r.tardiness_minutes or 0
        total_ot += r.overtime_minutes or 0
        if (r.tardiness_minutes or 0) > 0:
            emp_with_tard.add(r.employee_id)

    return {
        "period": {"start": str(date_start or ""), "end": str(date_end or "")},
        "total_records": len(records),
        "by_status": by_status,
        "total_tardiness_minutes": total_tard,
        "total_overtime_minutes": total_ot,
        "employees_with_tardiness": len(emp_with_tard),
    }


@router.get("/export")
def export_processed(
    date_start: Optional[date] = Query(None),
    date_end: Optional[date] = Query(None),
    employee_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    query = db.query(ProcessedAttendance, EmployeeModel).join(
        EmployeeModel, ProcessedAttendance.employee_id == EmployeeModel.id
    )
    if date_start:
        query = query.filter(ProcessedAttendance.date >= date_start)
    if date_end:
        query = query.filter(ProcessedAttendance.date <= date_end)
    if employee_id:
        query = query.filter(ProcessedAttendance.employee_id == employee_id)
    rows = query.order_by(ProcessedAttendance.date.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asistencia Procesada"

    headers = ["Empleado", "Fecha", "1ra Entrada", "Últ. Salida",
               "Hs Total", "Tard.(Min)", "Sal.Ant.(Min)", "HsEx.(Min)", "Estado", "Justificación"]
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = center

    for row_idx, (pa, emp) in enumerate(rows, 2):
        name = f"{emp.last_name or ''}, {emp.first_name or ''}".strip(", ") or emp.identifier
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=str(pa.date))
        ws.cell(row=row_idx, column=3, value=pa.first_in.strftime("%H:%M") if pa.first_in else "")
        ws.cell(row=row_idx, column=4, value=pa.last_out.strftime("%H:%M") if pa.last_out else "")
        ws.cell(row=row_idx, column=5, value=pa.total_hours)
        ws.cell(row=row_idx, column=6, value=pa.tardiness_minutes)
        ws.cell(row=row_idx, column=7, value=pa.early_departure_minutes)
        ws.cell(row=row_idx, column=8, value=pa.overtime_minutes)
        ws.cell(row=row_idx, column=9, value=pa.status)
        ws.cell(row=row_idx, column=10, value=pa.justification or "")

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(12, max_len + 2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"asistencia_{date_start or 'all'}_{date_end or 'all'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
