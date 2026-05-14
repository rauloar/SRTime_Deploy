from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
import io

from core.logging_config import get_logger
logger = get_logger(__name__)

from api.deps import get_db, get_current_user
from api.schemas import UserResponse, UserCreate, UserBulkImportRow
from models.employee import User as EmployeeModel
from models.user import AuthUser

router = APIRouter()

class UserUpdate(BaseModel):
    identifier: Optional[str] = None
    first_name: Optional[str] = None
    last_name: Optional[str] = None
    is_active: Optional[bool] = None
    shift_id: Optional[int] = None

@router.get("/", response_model=List[UserResponse])
def read_users(skip: int = 0, limit: int = 100, db: Session = Depends(get_db), current_user: AuthUser = Depends(get_current_user)):
    users = db.query(EmployeeModel).offset(skip).limit(limit).all()
    return users

@router.get("/{user_id}", response_model=UserResponse)
def read_user(user_id: int, db: Session = Depends(get_db), current_user: AuthUser = Depends(get_current_user)):
    user = db.query(EmployeeModel).filter(EmployeeModel.id == user_id).first()
    if user is None:
        raise HTTPException(status_code=404, detail="User not found")
    return user
    
@router.post("/", response_model=UserResponse)
def create_user(user: UserCreate, db: Session = Depends(get_db), current_user: AuthUser = Depends(get_current_user)):
    db_user = db.query(EmployeeModel).filter(EmployeeModel.identifier == user.identifier).first()
    if db_user:
        raise HTTPException(status_code=400, detail="Identifier already registered")
    new_employee = EmployeeModel(
        identifier=user.identifier,
        first_name=user.first_name,
        last_name=user.last_name,
        is_active=user.is_active,
        shift_id=user.shift_id
    )
    db.add(new_employee)
    db.commit()
    db.refresh(new_employee)
    logger.info(f"Employee created: {user.identifier}", extra={"user": current_user.username, "action": "create_employee", "detail": user.identifier})
    return new_employee

@router.put("/{user_id}", response_model=UserResponse)
def update_user(user_id: int, user: UserUpdate, db: Session = Depends(get_db), current_user: AuthUser = Depends(get_current_user)):
    db_user = db.query(EmployeeModel).filter(EmployeeModel.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")
    if user.identifier is not None:
        existing = db.query(EmployeeModel).filter(EmployeeModel.identifier == user.identifier, EmployeeModel.id != user_id).first()
        if existing:
            raise HTTPException(status_code=400, detail="Identifier already registered by another user")
        db_user.identifier = user.identifier
    if user.first_name is not None: db_user.first_name = user.first_name
    if user.last_name is not None: db_user.last_name = user.last_name
    if user.is_active is not None: db_user.is_active = user.is_active
    if user.shift_id is not None: db_user.shift_id = user.shift_id
    else: db_user.shift_id = None  # Allow removing shift
    db.commit()
    db.refresh(db_user)
    logger.info(f"Employee updated: {db_user.identifier}", extra={"user": current_user.username, "action": "update_employee", "detail": db_user.identifier})
    return db_user

@router.delete("/{user_id}")
def delete_user(user_id: int, db: Session = Depends(get_db), current_user: AuthUser = Depends(get_current_user)):
    db_user = db.query(EmployeeModel).filter(EmployeeModel.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")
    db.delete(db_user)
    db.commit()
    logger.info(f"Employee deleted: {db_user.identifier}", extra={"user": current_user.username, "action": "delete_employee", "detail": db_user.identifier})
    return {"ok": True, "message": f"User {db_user.identifier} deleted"}


@router.get("/export")
def export_users(
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from models.shift import Shift

    employees = db.query(EmployeeModel).order_by(EmployeeModel.last_name, EmployeeModel.first_name).all()
    shift_ids = {e.shift_id for e in employees if e.shift_id}
    shifts_by_id = {s.id: s.name for s in db.query(Shift).filter(Shift.id.in_(shift_ids)).all()} if shift_ids else {}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empleados"

    headers = ["Identificador", "Nombre", "Apellido", "Turno", "Activo"]
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = center

    for row_idx, emp in enumerate(employees, 2):
        ws.cell(row=row_idx, column=1, value=emp.identifier)
        ws.cell(row=row_idx, column=2, value=emp.first_name or "")
        ws.cell(row=row_idx, column=3, value=emp.last_name or "")
        ws.cell(row=row_idx, column=4, value=shifts_by_id.get(emp.shift_id, "Sin turno"))
        ws.cell(row=row_idx, column=5, value="Sí" if emp.is_active else "No")

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(12, max_len + 2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=empleados.xlsx"}
    )


@router.post("/bulk-import")
def bulk_import_users(
    users: List[UserBulkImportRow],
    db: Session = Depends(get_db),
    current_user: AuthUser = Depends(get_current_user)
):
    created = skipped = 0
    for u in users:
        if db.query(EmployeeModel).filter_by(identifier=u.identifier).first():
            skipped += 1
            continue
        employee = EmployeeModel(
            identifier=u.identifier,
            first_name=u.first_name,
            last_name=u.last_name,
            is_active=u.is_active,
            shift_id=u.shift_id,
        )
        db.add(employee)
        created += 1
    db.commit()
    logger.info(f"Bulk import: {created} created, {skipped} skipped", extra={"user": current_user.username, "action": "bulk_import"})
    return {"created": created, "skipped": skipped}
